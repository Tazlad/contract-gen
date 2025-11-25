#!/usr/bin/env python3
"""
Script to create the Word document template for employee contracts
This needs to be run once to generate the template file
"""

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_template():
    doc = Document()

    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(11)

    # ===== LETTER OF OFFER =====

    # Header
    p = doc.add_paragraph()
    p.add_run('Letter of offer\n').bold = True
    p.add_run('Unique Needs\n')
    p.add_run('3/322 Hobart road\n')
    p.add_run('Youngtown TAS 7249\n')
    p.add_run('ABN: 43618574031\n')

    doc.add_paragraph(f'Date: {{{{OfferDate}}}}')
    doc.add_paragraph(f'{{{{EmployeeName}}}}\n{{{{EmployeeAddress}}}}')

    doc.add_paragraph(f'Dear {{{{EmployeeName}}}},')

    p = doc.add_paragraph()
    p.add_run('Re: Offer of employment').bold = True

    doc.add_paragraph(f'I am pleased to offer you the position of {{{{PositionTitle}}}} with Unique Needs.')

    doc.add_paragraph('Please find attached your employment contract which provides the terms and conditions for this position. At a minimum, your terms and conditions are in accordance with the National Employment Standards and the Social, Community, Home Care and Disability Services Industry Award 2010 (MA000100). I've also attached the position description.')

    doc.add_paragraph('If you have any questions about your employment, please contact us on 03 63 888440. You can also contact the Fair Work Ombudsman (www.fairwork.gov.au) for help with minimum terms and conditions of employment.')

    doc.add_paragraph('To accept this offer and the attached terms and conditions, please sign and date this letter in the section below and sign the attached employment contract and return to us.')

    doc.add_paragraph('Congratulations – We look forward to you joining us.')

    doc.add_paragraph('Yours sincerely,')
    doc.add_paragraph('\n')
    doc.add_paragraph('Employer signature')
    doc.add_paragraph('Adam Hinds and Kellie Soule')
    doc.add_paragraph('Employer name and position')

    doc.add_paragraph('\n')

    # Acceptance section
    p = doc.add_paragraph()
    p.add_run('Acceptance of offer').bold = True

    p = doc.add_paragraph(f'I {{{{EmployeeName}}}}:')
    doc.add_paragraph('•\thave read and understand the terms and conditions in the attached employment contract')
    doc.add_paragraph('•\thave discussed any issues I have with these terms and conditions with my employer and they have considered and responded to any issues raised')
    doc.add_paragraph('•\thave received a copy of the contract and this letter of offer for my records.')

    doc.add_paragraph('\n\n')
    doc.add_paragraph('Employee signature')
    doc.add_paragraph('\nDate')

    # Page break
    doc.add_page_break()

    # ===== EMPLOYMENT CONTRACT =====

    p = doc.add_paragraph()
    p.add_run('Employment contract').bold = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(f'This is an employment contract dated {{{{ContractDate}}}}')

    doc.add_paragraph('Between Unique Needs, 3/322 Hobart road, Youngtown TAS 7249:')
    doc.add_paragraph('\n')
    doc.add_paragraph(f'and {{{{EmployeeName}}}}, {{{{EmployeeAddress}}}}:')

    doc.add_paragraph('\n')

    p = doc.add_paragraph()
    p.add_run('Employee details (for payroll and emergency):').bold = True

    doc.add_paragraph(f'Phone: {{{{EmployeePhone}}}}')
    doc.add_paragraph(f'Email: {{{{EmployeeEmail}}}}')
    doc.add_paragraph(f'Emergency Contact Name: {{{{EmergencyContactName}}}}')
    doc.add_paragraph(f'Emergency Contact Phone: {{{{EmergencyContactPhone}}}}')
    doc.add_paragraph(f'Tax File Number (TFN): {{{{EmployeeTFN}}}}')
    doc.add_paragraph(f'Super Fund Name: {{{{SuperFundName}}}}')
    doc.add_paragraph(f'Bank BSB: {{{{EmployeeBSB}}}}')
    doc.add_paragraph(f'Bank Account Number: {{{{EmployeeAccountNumber}}}}')

    doc.add_paragraph('\n')

    # Position section
    p = doc.add_paragraph()
    p.add_run('Position').bold = True

    doc.add_paragraph(f'You are being employed in the position of {{{{PositionTitle}}}}.')
    doc.add_paragraph('You are being employed on a casual basis, as required.')
    doc.add_paragraph('As a casual employee, we do not guarantee the days or hours you'll work, or how long you'll be employed for. We do not commit to providing you with work that will be continuing or indefinite.')
    doc.add_paragraph('We may choose to offer you work and you may accept or refuse our offer. You will be paid a casual loading.')
    doc.add_paragraph('In some circumstances, you may have the right to become a permanent employee (known as 'casual conversion').')

    # Employment dates
    p = doc.add_paragraph()
    p.add_run('Employment dates').bold = True
    doc.add_paragraph(f'Your start date will be {{{{StartDate}}}}.')

    # Workplace
    p = doc.add_paragraph()
    p.add_run('Workplace').bold = True
    doc.add_paragraph('You may also be required to work at other locations where reasonable.')

    # Duties
    p = doc.add_paragraph()
    p.add_run('Duties').bold = True
    doc.add_paragraph('You will perform the following duties as required:')

    doc.add_paragraph('\nPersonal Care', style='List Bullet')
    doc.add_paragraph('Assist with bathing, dressing, grooming, and other routine tasks while respecting each client's dignity, privacy, and autonomy.')

    doc.add_paragraph('\nHealth Monitoring', style='List Bullet')
    doc.add_paragraph('Observe changes in physical, mental, or emotional well-being, and report concerns to supervisors or healthcare professionals.')

    doc.add_paragraph('\nMobility & Transfer Support', style='List Bullet')
    doc.add_paragraph('Use approved lifting and transfer techniques, ensuring safe practices for both the client and yourself.')

    doc.add_paragraph('\nMedication Assistance', style='List Bullet')
    doc.add_paragraph('Remind or administer medication if authorized and endorsed, keeping accurate records of dosages, schedules, and any notable side effects.')

    doc.add_paragraph('\nHousehold Tasks', style='List Bullet')
    doc.add_paragraph('Help clients maintain clean, safe living environments with light housekeeping, laundry, and meal preparation, encouraging active participation where possible.')

    doc.add_paragraph('\nSkill Development', style='List Bullet')
    doc.add_paragraph('Promote independence by teaching and reinforcing daily living skills like cooking, shopping, and budgeting, following individualized care plans.')

    doc.add_paragraph('\nCommunity Participation', style='List Bullet')
    doc.add_paragraph('Accompany or transport clients to appointments, recreational outings, or social events, supporting community engagement and inclusion.')

    doc.add_paragraph('\nEmotional & Behavioral Support', style='List Bullet')
    doc.add_paragraph('Provide a calm, empathetic presence and follow behavior support strategies, upholding professional boundaries to ensure a therapeutic relationship.')

    doc.add_paragraph('\nDocumentation & Reporting', style='List Bullet')
    doc.add_paragraph('Maintain clear, accurate notes on daily activities, incidents, and client progress, while safeguarding confidentiality and privacy.')

    doc.add_paragraph('\nCollaboration & Compliance', style='List Bullet')
    doc.add_paragraph('Coordinate with families, healthcare providers, and team members. Adhere to Unique Needs' policies, safety procedures, and legal requirements, including mandatory reporting and respect for professional boundaries.')

    doc.add_paragraph('\nWe may also assign you other duties, where reasonable for your position, training and experience.')

    # Employment terms and conditions
    p = doc.add_paragraph()
    p.add_run('Employment terms and conditions').bold = True
    doc.add_paragraph('Your employment terms and conditions are those set out in this contract, the Social, Community, Home Care and Disability Services Industry Award 2010 (MA000100) and applicable legislation. This includes, the National Employment Standards in the Fair Work Act 2009.')
    doc.add_paragraph('You can check the minimum award entitlements for your classification level with the Fair Work Ombudsman's Pay and Conditions Tool.')

    # Hours of work
    p = doc.add_paragraph()
    p.add_run('Hours of work').bold = True
    doc.add_paragraph('As a casual employee, your hours may vary depending on business needs.')

    p = doc.add_paragraph()
    p.add_run('Shift work hours').bold = True
    doc.add_paragraph('At the start of your employment, you will be rostered to work the following shift type in accordance with the award:')
    doc.add_paragraph('shifts may include day, afternoon/evening, weekend, overnight (active/sleepover), public holiday, and on-call. Penalty rates apply where required.')
    doc.add_paragraph('The shifts you work may be changed later in accordance with award rules about changing hours of work, rosters and consultation.')

    p = doc.add_paragraph()
    p.add_run('Breaks').bold = True
    doc.add_paragraph('Depending on the number of hours you work in a shift, you may be entitled to meal breaks. The award sets out:')
    doc.add_paragraph('•\tthe length of the breaks')
    doc.add_paragraph('•\twhen they need to be taken')
    doc.add_paragraph('•\tthe rules about payment.')
    doc.add_paragraph('You may also be entitled to rest breaks.')

    p = doc.add_paragraph()
    p.add_run('On call or standby').bold = True
    doc.add_paragraph('We may require you to be on call (on stand-by) for work outside your normal hours.')
    doc.add_paragraph('Your on call or stand-by conditions will be in accordance with the award, including:')
    doc.add_paragraph('•\trostering to be on call')
    doc.add_paragraph('•\tpay rates or allowances when on call')
    doc.add_paragraph('•\tpay rates if you are called out to work while on call.')

    # Pay and allowances
    p = doc.add_paragraph()
    p.add_run('Pay and allowances').bold = True

    p = doc.add_paragraph()
    p.add_run('Pay rate').bold = True
    doc.add_paragraph(f'You will be paid {{{{PayRate}}}} per hour. This pay rate includes casual loading at the percentage set out in your award. This loading is paid instead of entitlements that apply to permanent employees like paid personal leave and annual leave.')
    doc.add_paragraph('This pay rate does not include superannuation, we'll pay this separately.')

    p = doc.add_paragraph()
    p.add_run('Payment method').bold = True
    doc.add_paragraph('We will pay you fortnightly into your nominated bank account.')

    p = doc.add_paragraph()
    p.add_run('Penalty rates and overtime').bold = True
    doc.add_paragraph('You may be entitled to overtime rates under your award if you work:')
    doc.add_paragraph('•\tmore than your ordinary hours of work')
    doc.add_paragraph('•\toutside the spread of ordinary hours.')
    doc.add_paragraph('You may be entitled to penalty rates or shift loadings according to your award if you work:')
    doc.add_paragraph('•\ton a weekend')
    doc.add_paragraph('•\ton a public holiday')
    doc.add_paragraph('•\tlate night or early morning shifts.')

    p = doc.add_paragraph()
    p.add_run('Superannuation').bold = True
    doc.add_paragraph('If you are eligible for the super guarantee (SG), we will pay the contributions on your behalf in accordance with legislation and your award. We will pay contributions into a super fund of your choice.')
    doc.add_paragraph('If you do not tell us your choice of fund, we may need to contact the ATO to find out if you have a 'stapled' super fund to make your SG contributions into.')
    doc.add_paragraph('If you do not tell us your choice of fund and the ATO confirms you don't have a stapled super fund, we will pay your SG contributions to our default fund or another fund that meets the choice of fund rules.')

    p = doc.add_paragraph()
    p.add_run('Annual bonus').bold = True
    doc.add_paragraph('We may pay you an annual bonus at our discretion, taking into consideration your performance and other relevant business factors.')
    doc.add_paragraph('There is no guarantee that you will get a bonus, and you will not be eligible for one after your employment with us ends.')

    p = doc.add_paragraph()
    p.add_run('Annual pay review').bold = True
    doc.add_paragraph('We will review your pay annually to determine whether you are eligible for an increase, taking into consideration:')
    doc.add_paragraph('•\tyour performance')
    doc.add_paragraph('•\tthe business's financial position.')
    doc.add_paragraph('Any increase in your pay, above your award entitlements, is our decision.')

    p = doc.add_paragraph()
    p.add_run('Allowances').bold = True
    doc.add_paragraph('You may be entitled to allowances in accordance with your award, for example, if you:')
    doc.add_paragraph('•\tdo certain tasks or have a particular skill')
    doc.add_paragraph('•\thave to use your own tools at work')
    doc.add_paragraph('•\twork in particular conditions, environments or in remote locations.')

    # Leave
    p = doc.add_paragraph()
    p.add_run('Leave').bold = True

    p = doc.add_paragraph()
    p.add_run('Carer's leave').bold = True
    doc.add_paragraph('You are entitled to 2 days of unpaid carer's leave (in accordance with the National Employment Standards). This is available each time a member of your immediate family or household needs your care or support because of:')
    doc.add_paragraph('•\tpersonal injury')
    doc.add_paragraph('•\tpersonal illness')
    doc.add_paragraph('•\tan unexpected emergency.')
    doc.add_paragraph('You must give us notice as soon as possible to take carer's leave. We may also require evidence (such as a medical certificate).')

    p = doc.add_paragraph()
    p.add_run('Compassionate leave').bold = True
    doc.add_paragraph('You are entitled to 2 days unpaid compassionate leave (in accordance with the National Employment Standards) each time:')
    doc.add_paragraph('•\ta member of your immediate family or household dies, or has a life threatening illness or injury')
    doc.add_paragraph('•\ta child is stillborn, that would have been a member of your immediate family or household if born alive')
    doc.add_paragraph('•\tyou or your current spouse or de facto partner has a miscarriage.')

    p = doc.add_paragraph()
    p.add_run('Community service leave').bold = True
    doc.add_paragraph('You are entitled to community service leave (in accordance with the National Employment Standards) for certain activities such as:')
    doc.add_paragraph('•\tvoluntary emergency management activities')
    doc.add_paragraph('•\tjury duty and jury selection.')
    doc.add_paragraph('You must give us:')
    doc.add_paragraph('•\tnotice of your leave as soon as possible')
    doc.add_paragraph('•\tdetails of the period, or expected period, that you will be away from work.')
    doc.add_paragraph('We may ask you to provide evidence that you require community service leave.')

    p = doc.add_paragraph()
    p.add_run('Family and domestic violence leave').bold = True
    doc.add_paragraph('You are entitled to 10 days of paid family and domestic violence leave (in accordance with the National Employment Standards) each 12 month period.')

    p = doc.add_paragraph()
    p.add_run('Long service leave').bold = True
    doc.add_paragraph('You may be entitled to long service leave after working with us for a specific period of time in accordance with relevant long service leave legislation.')

    p = doc.add_paragraph()
    p.add_run('Parental leave').bold = True
    doc.add_paragraph('You may be entitled to 12 months of unpaid parental leave if you:')
    doc.add_paragraph('•\thave worked for us on a regular and systematic basis for 12 months or more')
    doc.add_paragraph('•\thad a reasonable expectation of continuing this regular work had it not been for the birth or adoption of a child.')
    doc.add_paragraph('You may also:')
    doc.add_paragraph('•\trequest up to an extra 12 months of unpaid leave')
    doc.add_paragraph('•\tbe entitled to Parental Leave Pay from the Australian Government, administered by Services Australia.')

    p = doc.add_paragraph()
    p.add_run('Public holidays').bold = True
    doc.add_paragraph('You have a right to be absent from work on a public holiday.')
    doc.add_paragraph('We may ask you to work on a public holiday, but as you are a casual employee you may refuse our offer of work.')
    doc.add_paragraph('If you work on a public holiday, you are entitled to any additional entitlements set out under your award, such as public holiday penalty rates.')

    # Obligations
    p = doc.add_paragraph()
    p.add_run('Obligations').bold = True

    p = doc.add_paragraph()
    p.add_run('Employee obligations').bold = True
    doc.add_paragraph('As an employee of our business we expect you to:')
    doc.add_paragraph('•\tcarry out your duties to the best of your ability')
    doc.add_paragraph('•\tact honestly and in the best interests of the business')
    doc.add_paragraph('•\tcomply with our business policies and procedures which we will make available to you (but do not form part of this contract)')
    doc.add_paragraph('•\tcomply with any other lawful and reasonable directions we provide.')

    p = doc.add_paragraph()
    p.add_run('Conflict of interest').bold = True
    doc.add_paragraph('While employed with us, you must get our written agreement before working for other employers or doing activities that may conflict with the interests of our business.')

    p = doc.add_paragraph()
    p.add_run('Confidentiality').bold = True
    doc.add_paragraph('You agree not to use or disclose confidential information relating to the business. This includes while you are employed by us and after your employment ends.')
    doc.add_paragraph('Confidential information – including trade secrets, pricing structures, documents you create while employed with us, and information on our clients and suppliers – is our property.')
    doc.add_paragraph('There are exceptions if:')
    doc.add_paragraph('•\twe have given you our consent')
    doc.add_paragraph('•\tyou are using the information appropriately to do your work for us')
    doc.add_paragraph('•\tthe information is already publicly available')
    doc.add_paragraph('•\tthe information is required by law.')

    p = doc.add_paragraph()
    p.add_run('Intellectual property').bold = True
    doc.add_paragraph('Anything you invent, develop or create in the course of your work with us, remains our intellectual property. You must tell us about these works immediately.')
    doc.add_paragraph('This includes:')
    doc.add_paragraph('•\tdesigns')
    doc.add_paragraph('•\tlogos')
    doc.add_paragraph('•\tbusiness and domain names')
    doc.add_paragraph('•\tcopyright')
    doc.add_paragraph('•\ttrade marks')
    doc.add_paragraph('•\tpatents.')
    doc.add_paragraph('You must not use or reproduce any intellectual property owned by us without our consent. This includes after your employment ends with us.')

    # Ending employment
    p = doc.add_paragraph()
    p.add_run('Ending employment').bold = True

    p = doc.add_paragraph()
    p.add_run('Notice').bold = True
    doc.add_paragraph('As a casual employee, if:')
    doc.add_paragraph('•\twe end your employment, we do not have to give you notice')
    doc.add_paragraph('•\tyou resign, you do not have to give us notice.')

    p = doc.add_paragraph()
    p.add_run('Misconduct').bold = True
    doc.add_paragraph('We may terminate your employment without notice, or payment in lieu of notice, if you engage in serious misconduct.')
    doc.add_paragraph('Serious misconduct is when an employee:')
    doc.add_paragraph('•\tcauses serious and imminent risk to the health and safety of another person or to the reputation, viability or profits of their employer's business, or')
    doc.add_paragraph('•\twilfully or deliberately behaves in a way that's inconsistent with continuing their employment.')
    doc.add_paragraph('Examples of serious misconduct include:')
    doc.add_paragraph('•\ttheft')
    doc.add_paragraph('•\tfraud')
    doc.add_paragraph('•\tviolence/assault')
    doc.add_paragraph('•\tsexual harassment')
    doc.add_paragraph('•\tserious breaches of health and safety requirements')
    doc.add_paragraph('•\tbeing drunk or affected by drugs at work')
    doc.add_paragraph('•\trefusing to carry out work duties.')

    # Signature section
    doc.add_page_break()

    p = doc.add_paragraph()
    p.add_run('SIGNED').bold = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    p.add_run('Employer – partnership or association').bold = True

    doc.add_paragraph('SIGNED for and on behalf of A Hinds & K Soule / Unique Needs, ABN: 43618574031 by a duly authorised representative who warrants that they have the authority to sign this contract on behalf of A Hinds & K Soule / Unique Needs.')

    doc.add_paragraph('\nAdam Hinds / Kellie Soule\t\t')
    doc.add_paragraph('Name of authorised representative (print)\t-\tName of witness (print)')
    doc.add_paragraph('\n\t\t')
    doc.add_paragraph('Signature of authorised representative\t\tSignature of witness')
    doc.add_paragraph(f'{{{{ContractDate}}}}\t\t{{{{ContractDate}}}}')
    doc.add_paragraph('Date\t\tDate')

    doc.add_paragraph('\n')

    p = doc.add_paragraph()
    p.add_run('Employee').bold = True

    doc.add_paragraph('I understand and agree to the terms and conditions of employment set out in this contract.')

    doc.add_paragraph(f'\n{{{{EmployeeName}}}}')
    doc.add_paragraph('Name (print)')
    doc.add_paragraph('\n\nSignature')
    doc.add_paragraph('\n\nDate')

    # Save the document
    doc.save('templates/contract_template.docx')
    print('Template created successfully at templates/contract_template.docx')

if __name__ == '__main__':
    create_template()
